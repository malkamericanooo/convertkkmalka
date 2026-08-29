import openpyxl
wb = openpyxl.load_workbook("/Users/malkaarif/mabuun-konverter/output/REKAP MABUUN  converted.xlsx")
ws = wb.active
green=0; kk=0; mis=0
for r in range(1, ws.max_row+1):
    v4=ws.cell(r,4).value
    if v4 in (None,''): continue
    f=ws.cell(r,4).fill
    g=f.patternType=='solid' and f.fgColor.rgb=='FFCCFFCC'
    hub=str(ws.cell(r,5).value).upper()=='KK'
    if g: green+=1
    if hub: kk+=1
    if g!=hub: mis+=1
print('CONVERTED: green:', green, '| KK:', kk, '| mismatches:', mis)
print('r19:', repr(ws.cell(19,1).value), repr(ws.cell(19,2).value), repr(ws.cell(19,3).value), repr(ws.cell(19,4).value))
print('r20:', repr(ws.cell(20,1).value), repr(ws.cell(20,2).value), repr(ws.cell(20,4).value))
txt = [ws.cell(r,1).value for r in range(ws.max_row-60, ws.max_row+1) if ws.cell(r,1).value]
print('tail A-col:', txt[-12:])
print()
wb2 = openpyxl.load_workbook("/Users/malkaarif/mabuun-konverter/output/DESA MABUUN.xlsx")
w2 = wb2.active
print('SUMMARY:')
for r in range(1, 22):
    vals = [w2.cell(r,c).value for c in range(1,6)]
    if any(vals): print(r, ' | '.join(str(v)[:22] if v else '' for v in vals))
