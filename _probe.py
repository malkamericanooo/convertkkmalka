import openpyxl
p = "/Users/malkaarif/Downloads/mabuun project 3/MABU'UN Rekapitulasi Data Keluarga (29) (1).xlsx"
wb = openpyxl.load_workbook(p)
ws = wb['Sheet1']
print('max_row:', ws.max_row)
for r in range(ws.max_row-30, ws.max_row+1):
    vals = []
    for c in range(1, 10):
        v = ws.cell(r,c).value
        if v is not None: vals.append(f'{c}:{str(v)[:40]}')
    if vals: print(r, ' || '.join(vals))
